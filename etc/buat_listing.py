"""
Membuat dokumen listing kode Python TGS 3.
Output: TGS_3_Kelompok_1_Klasifikasi_Email_Spam_NRP_Nama_Listing_Code.docx

Berisi daftar isi file, instruksi menjalankan, dan listing lengkap:
  - buat_dataset.py
  - klasifikasi_spam.py
"""

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(BASE, "TGS_3_Kelompok_1_Klasifikasi_Email_Spam_NRP_Nama_Listing_Code.docx")
FILES = ["buat_dataset.py", "klasifikasi_spam.py"]


def judul(doc, teks, ukuran=14, center=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(teks)
    r.bold = True
    r.font.size = Pt(ukuran)
    r.font.name = "Times New Roman"
    return p


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)

    # footer nomor halaman
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)

    judul(doc, "TGS 3 - DATA MINING & INFORMATION RETRIEVAL (IFB-307)", center=True)
    judul(doc, "LISTING KODE PROGRAM - KLASIFIKASI EMAIL SPAM", ukuran=16, center=True)
    judul(doc, "Kelompok 1 - [Nama Anggota (NRP)]", ukuran=12, center=True)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    judul(doc, "A. RINGKASAN PROGRAM", ukuran=13)
    for baris in [
        "Program utama : klasifikasi_spam.py",
        "Pembuat data   : buat_dataset.py (dataset dummy, seed = 42)",
        "Data           : spam_dummy.csv (240 email, 5 fitur + label 0/1)",
        "Library        : scikit-learn, pandas, numpy, matplotlib (Python 3)",
        "",
        "Perintah menjalankan:",
        "  [KLIK GANDA] Jalankan Program.bat              # cara termudah dari File Explorer",
        "                                                 # (latih + evaluasi + menu, jendela tetap terbuka)",
        "  python buat_dataset.py                          # membuat ulang dataset dummy",
        "  python klasifikasi_spam.py                      # latih + evaluasi 3 model (+ menu interaktif)",
        "  python klasifikasi_spam.py --prediksi 6 4 4 12 3  # uji 1 email baru",
        "  python klasifikasi_spam.py --data data_lain.csv # pakai file CSV lain",
        "  python klasifikasi_spam.py --menu               # menu interaktif",
        "",
        "Model: Naive Bayes (utama), Logistic Regression (pembanding), "
        "Decision Tree (visual).",
        "Letak klasifier: fungsi buat_model() dan latih_dan_evaluasi() "
        "(perintah model.fit).",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(baris)
        r.font.name = "Consolas"
        r.font.size = Pt(10.5)

    for nama_file in FILES:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        judul(doc, f"B. LISTING {nama_file.upper()}", ukuran=13)
        with open(os.path.join(BASE, nama_file), encoding="utf-8") as f:
            for i, baris in enumerate(f, start=1):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                r = p.add_run(f"{i:>4} | {baris.rstrip()}")
                r.font.name = "Consolas"
                r.font.size = Pt(9)

    doc.save(OUT)
    print(f"Listing dibuat: {OUT}")


if __name__ == "__main__":
    main()
